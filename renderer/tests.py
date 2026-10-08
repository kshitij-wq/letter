import io
import json

from django.test import Client, SimpleTestCase
from docx import Document

from .engine import format_date, indian_comma, int_comma, render_docx, RenderError


def make_docx(paragraphs, table=None):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, text in enumerate(row):
                t.cell(r, c).text = text
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class FilterTests(SimpleTestCase):
    def test_indian_comma(self):
        self.assertEqual(indian_comma(1200000), "12,00,000")
        self.assertEqual(indian_comma("1200000"), "12,00,000")
        self.assertEqual(indian_comma("12,00,000"), "12,00,000")
        self.assertEqual(indian_comma(4900.0), "4,900")
        self.assertEqual(indian_comma(-45000), "-45,000")
        self.assertEqual(indian_comma("N/A"), "N/A")
        self.assertEqual(indian_comma(""), "")

    def test_int_comma(self):
        self.assertEqual(int_comma(1200000), "1,200,000")

    def test_format_date(self):
        self.assertEqual(format_date("01-04-2026"), "April 01, 2026")
        self.assertEqual(format_date("2026-04-01", "%Y-%m-%d", "%d %B %Y"), "01 April 2026")
        self.assertEqual(format_date("not a date"), "not a date")


class RenderTests(SimpleTestCase):
    def test_render_house_style_and_rows(self):
        tpl = make_docx(
            ["Dear {{ first_name }},", "CTC: Rs. {{ '{:.0f}'.format(ctc | float) | indian_comma }}",
             'Dept: {% if dept == "R\\u0026D" %}research{% else %}other{% endif %}'],
            table=[["Basic", "{% if (basic | default(0) | float) != 0 %}{{ '{:.0f}'.format(basic | float) | indian_comma }}{% endif %}"]],
        )
        out = render_docx(tpl, {"first_name": "Asha", "ctc": "1200000", "dept": "R&D", "basic": 45000}, autoescape=True)
        doc = Document(io.BytesIO(out))
        texts = [p.text for p in doc.paragraphs]
        self.assertIn("Dear Asha,", texts)
        self.assertIn("CTC: Rs. 12,00,000", texts)
        self.assertIn("Dept: research", texts)
        self.assertEqual(doc.tables[0].cell(0, 1).text, "45,000")

    def test_ampersand_in_data_is_escaped(self):
        tpl = make_docx(["Dept: {{ dept }}"])
        out = render_docx(tpl, {"dept": "R&D <APAC>"}, autoescape=True)
        self.assertEqual(Document(io.BytesIO(out)).paragraphs[0].text, "Dept: R&D <APAC>")

    def test_missing_field_with_float_is_explained(self):
        tpl = make_docx(["{{ old_basic | float }}"])
        with self.assertRaises(RenderError) as cm:
            render_docx(tpl, {})
        self.assertIn("old_basic", cm.exception.message)

    def test_syntax_error_reports_tag(self):
        tpl = make_docx(["{% if x %}open but never closed"])
        with self.assertRaises(RenderError) as cm:
            render_docx(tpl, {"x": 1})
        self.assertTrue(cm.exception.message)

    def test_sandbox_blocks_python_internals(self):
        tpl = make_docx(["{{ ''.__class__.__mro__ }}"])
        with self.assertRaises(RenderError):
            render_docx(tpl, {})


class ApiTests(SimpleTestCase):
    def setUp(self):
        self.client = Client()

    def test_index_serves_the_tool(self):
        resp = self.client.get("/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn(b"Letter Tagger", resp.content)
        self.assertIn("csrftoken", resp.cookies)

    def test_health(self):
        data = self.client.get("/api/health").json()
        self.assertEqual(data["app"], "Letter Studio")
        self.assertTrue(data["validate"])

    def test_render_json(self):
        tpl = make_docx(["Hello {{ first_name }}"])
        resp = self.client.post("/api/render", {
            "template": self._file(tpl),
            "context": json.dumps({"first_name": "Ravi"}),
            "format": "json",
        })
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertIn("Hello Ravi", resp.json()["text"]["body"])

    def test_render_error_is_422(self):
        tpl = make_docx(["{{ x | float }}"])
        resp = self.client.post("/api/render", {"template": self._file(tpl), "context": "{}", "format": "docx"})
        self.assertEqual(resp.status_code, 422)
        self.assertIn("error", resp.json())

    def test_bad_requests_get_plain_json_errors(self):
        tpl = make_docx(["Hi"])
        cases = [
            ({}, 400),                                                            # no letter
            ({"template": self._file(tpl, "letter.doc")}, 400),                   # not .docx
            ({"template": self._file(tpl), "context": "{not json"}, 400),         # broken JSON
            ({"template": self._file(tpl), "context": "[1, 2]"}, 400),            # JSON but not values
            ({"template": self._file(tpl), "format": "pdf"}, 400),                # no PDFs any more
            ({"template": self._file(b"not a zip")}, 422),                        # not a Word file
        ]
        for data, status in cases:
            resp = self.client.post("/api/render", data)
            self.assertEqual(resp.status_code, status, (data.keys(), resp.content[:200]))
            self.assertIn("error", resp.json())
        self.assertEqual(self.client.get("/api/render").status_code, 405)

    def test_too_large_is_413(self):
        from django.test import override_settings
        with override_settings(MAX_UPLOAD_MB=0):
            resp = self.client.post("/api/render", {"template": self._file(make_docx(["Hi"])), "format": "json"})
        self.assertEqual(resp.status_code, 413)

    def test_unexpected_failure_is_json_500(self):
        from unittest import mock
        with mock.patch("renderer.views.render_docx", side_effect=RuntimeError("boom")), self.assertLogs("renderer", "ERROR"):
            resp = self.client.post("/api/render", {"template": self._file(make_docx(["Hi"])), "format": "json"})
        self.assertEqual(resp.status_code, 500)
        self.assertEqual(resp.json()["kind"], "server")

    def test_csrf_is_required_from_a_browser(self):
        strict = Client(enforce_csrf_checks=True)
        tpl = make_docx(["Hello {{ first_name }}"])
        self.assertEqual(strict.post("/api/render", {"template": self._file(tpl), "format": "json"}).status_code, 403)
        strict.get("/")
        token = strict.cookies["csrftoken"].value
        resp = strict.post("/api/render", {"template": self._file(tpl), "format": "json", "context": '{"first_name": "Ravi"}'}, HTTP_X_CSRFTOKEN=token)
        self.assertEqual(resp.status_code, 200)

    def test_headers_and_footers_are_filled(self):
        doc = Document()
        doc.add_paragraph("Body {{ name }}")
        doc.sections[0].header.paragraphs[0].text = "Header {{ name }}"
        doc.sections[0].footer.paragraphs[0].text = "Footer {{ code }}"
        buf = io.BytesIO(); doc.save(buf)
        resp = self.client.post("/api/render", {"template": self._file(buf.getvalue()), "format": "json", "context": '{"name": "Asha", "code": "E1"}'})
        text = resp.json()["text"]
        self.assertIn("Header Asha", text["headers"])
        self.assertIn("Footer E1", text["footers"])

    def test_download_name_is_safe(self):
        resp = self.client.post("/api/render", {"template": self._file(make_docx(["Hi"])), "format": "docx", "name": 'a"b\r\nX: y/../c.docx'})
        self.assertEqual(resp.status_code, 200)
        self.assertNotIn('"b', resp["Content-Disposition"])
        self.assertNotIn("\n", resp["Content-Disposition"])

    @staticmethod
    def _file(data, name="letter.docx"):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return SimpleUploadedFile(name, data, content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


class ValidateTests(SimpleTestCase):
    def post(self, paragraphs, employees=None, **extra):
        data = {"template": ApiTests._file(make_docx(paragraphs))}
        if employees is not None:
            data["employees"] = json.dumps(employees)
        data.update(extra)
        return self.client.post("/api/validate", data)

    def test_clean_template(self):
        resp = self.post(["Dear {{ first_name }},", "{% if band == \"A\" %}Top{% endif %}"])
        data = resp.json()
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(data["ok"])
        self.assertEqual(data["variables"], ["band", "first_name"])

    def test_syntax_errors_point_at_the_paragraph(self):
        cases = {
            "{{ ctc | floot }}": "floot",
            "{% if band = 1 %}x{% endif %}": "==",
            '{% if band == "A %}x{% endif %}': "quote",
            "{% endif %}": "without",
            "{% iff x %}": "isn't a Jinja tag",
        }
        for tag, words in cases.items():
            data = self.post(["Hello", "Second line", tag]).json()
            self.assertFalse(data["ok"], tag)
            err = data["syntax"][0]
            self.assertIn(words, err["friendly"], tag)
            self.assertEqual(err["para_index"], 2, tag)
            self.assertEqual(err["part_label"], "Body")

    def test_unclosed_if_has_no_misleading_place(self):
        err = self.post(["{% if x %}", "never closed"]).json()["syntax"][0]
        self.assertIn("never closed", err["friendly"])
        self.assertIsNone(err["para_index"])

    def test_header_errors_name_the_header(self):
        doc = Document()
        doc.add_paragraph("Body")
        doc.sections[0].header.paragraphs[0].text = "{{ name | floot }}"
        buf = io.BytesIO(); doc.save(buf)
        data = self.client.post("/api/validate", {"template": ApiTests._file(buf.getvalue())}).json()
        self.assertEqual(data["syntax"][0]["part_label"], "Header")

    def test_employees_that_break_the_letter_are_listed(self):
        emps = [{"label": "E1", "values": {"old": "5", "dept": "HR"}},
                {"label": "E2", "values": {"dept": "R&D"}}]
        data = self.post(["{{ old | float }}", "Dept {{ dept }}"], emps).json()
        self.assertEqual(data["employees"]["tested"], 2)
        failed = data["employees"]["failed"]
        self.assertEqual([f["label"] for f in failed], ["E2"])
        self.assertEqual(failed[0]["name"], "old")
        self.assertEqual(failed[0]["para_index"], 0)

    def test_unescaped_ampersand_is_a_broken_file(self):
        data = self.post(["Dept {{ dept }}"], [{"label": "E1", "values": {"dept": "R&D"}}], autoescape="0").json()
        self.assertEqual(data["employees"]["failed"][0]["kind"], "BrokenFile")
        data = self.post(["Dept {{ dept }}"], [{"label": "E1", "values": {"dept": "R&D"}}], autoescape="1").json()
        self.assertTrue(data["ok"])

    def test_bad_input(self):
        self.assertEqual(self.client.post("/api/validate", {}).status_code, 400)
        bad = self.client.post("/api/validate", {"template": ApiTests._file(make_docx(["Hi"])), "employees": "{nope"})
        self.assertEqual(bad.status_code, 400)
        resp = self.client.post("/api/validate", {"template": ApiTests._file(b"not a zip")})
        self.assertEqual(resp.status_code, 422)
        self.assertEqual(self.client.get("/api/validate").status_code, 405)
