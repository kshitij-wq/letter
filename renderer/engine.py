"""Fill a tagged .docx with one employee's values using real Jinja (docxtpl).

The filters below (indian_comma, int_comma, format_date, currentDatetime) copy how the
Letter Tagger preview assumes CompUp behaves. If CompUp's own filters differ, change
them here; everything else in the app follows.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from django.conf import settings
from docxtpl import DocxTemplate
from jinja2 import TemplateError, TemplateSyntaxError, Undefined, UndefinedError
from jinja2.sandbox import SandboxedEnvironment

NUMBER_RE = re.compile(r"[-+]?\d+(\.\d+)?")


class RenderError(Exception):
    """A template that Jinja can't fill, explained for people who write letters."""

    def __init__(self, message, kind="error", tag="", context=None, part=""):
        super().__init__(message)
        self.message = message
        self.kind = kind
        self.tag = tag
        self.context = context or []
        self.part = part

    def as_dict(self):
        return {"error": self.message, "kind": self.kind, "tag": self.tag, "context": self.context, "part": self.part}


# ---------------------------------------------------------------- filters

def _to_number(value):
    if isinstance(value, bool) or value is None or isinstance(value, Undefined):
        return None
    if isinstance(value, (int, float, Decimal)):
        return value
    text = str(value).strip().replace(",", "")
    if NUMBER_RE.fullmatch(text):
        return float(text) if "." in text else int(text)
    return None


def _group(number, indian):
    negative = number < 0
    number = abs(number)
    if isinstance(number, float) and number.is_integer():
        number = int(number)
    if isinstance(number, Decimal):
        text = format(number, "f")
    else:
        text = str(number) if isinstance(number, int) else repr(number)
    whole, _, decimals = text.partition(".")
    if indian:
        last3, rest = whole[-3:], whole[:-3]
        if rest:
            rest = re.sub(r"(\d)(?=(\d{2})+$)", r"\1,", rest) + ","
        whole = rest + last3
    else:
        whole = "{:,}".format(int(whole))
    return ("-" if negative else "") + whole + ("." + decimals if decimals else "")


def indian_comma(value):
    """12,00,000 style grouping. Text that isn't a number is printed as it is."""
    if value is None or value == "" or isinstance(value, Undefined):
        return ""
    number = _to_number(value)
    return str(value) if number is None else _group(number, indian=True)


def int_comma(value):
    """1,200,000 style grouping."""
    if value is None or value == "" or isinstance(value, Undefined):
        return ""
    number = _to_number(value)
    return str(value) if number is None else _group(number, indian=False)


def format_date(value, in_format="%d-%m-%Y", out_format="%B %d, %Y"):
    """Reformat a date. Values that don't match in_format are printed as they are."""
    if value is None or value == "" or isinstance(value, Undefined):
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime(out_format)
    try:
        return datetime.strptime(str(value).strip(), in_format).strftime(out_format)
    except ValueError:
        return str(value)


def current_datetime(fmt="%B %d, %Y"):
    return datetime.now(ZoneInfo(settings.TIME_ZONE)).strftime(fmt)


def make_env(autoescape=True):
    # Sandboxed: templates come from uploads, so they must not reach Python internals.
    env = SandboxedEnvironment(autoescape=autoescape, undefined=Undefined)
    env.filters.update({
        "indian_comma": indian_comma,
        "int_comma": int_comma,
        "format_date": format_date,
    })
    env.globals.update({"currentDatetime": current_datetime})
    return env


# ---------------------------------------------------------------- render

def _friendly(exc):
    text = str(exc)
    if isinstance(exc, UndefinedError):
        m = re.search(r"'([^']+)' is undefined", text)
        if m:
            return f"“{m.group(1)}” isn't in the data for this employee, and the tag uses it in a way that needs a value (for example with | float). Add the column, or write {m.group(1)} | default(0)."
        m = re.search(r"has no attribute '([^']+)'", text)
        if m:
            return f"“{m.group(1)}” is missing for this employee (for example custom_schema.{m.group(1)}), and the tag needs a value there. Add the column, or use | default(…)."
        return f"A value is missing for this employee: {text}"
    if isinstance(exc, TemplateSyntaxError):
        m = re.search(r"No filter named '([^']+)'", text)
        if m:
            return f"“| {m.group(1)}” isn't a filter Jinja knows. Check the spelling (for example | float, | int, | indian_comma, | lower)."
        m = re.search(r"Encountered unknown tag '(\w+)'", text)
        if m and m.group(1) in ("endif", "else", "elif", "endfor"):
            return f"There is a {{% {m.group(1)} %}} without the {{% if %}} (or {{% for %}}) it belongs to above it."
        if m:
            return f"“{{% {m.group(1)} %}}” isn't a Jinja tag. Check the spelling (if, elif, else, endif, for, endfor, set)."
        if "Unexpected end of template" in text:
            m = re.search(r"innermost block that needs to be closed is '(\w+)'", text)
            return f"An {{% {m.group(1) if m else 'if'} %}} is never closed. Add {{% end{m.group(1) if m else 'if'} %}} where it should stop."
        if "got '='" in text:
            return "Use == to compare inside a condition, e.g. {% if band == \"A\" %}. A single = only works in {% set %}."
        if "expected token 'end of statement block'" in text or "expected token 'end of print statement'" in text:
            got = re.search(r"got '([^']+)'", text)
            return f"A tag has a typing mistake near “{got.group(1)}”: a missing quote, bracket or | , or two words without and / or between them." if got else f"A tag has a typing mistake: {text}"
        if "Unexpected end of template" in text or "endif" in text and "Encountered unknown tag" in text:
            return f"An if or for block isn't closed properly: {text}"
        return f"A tag has a typing mistake: {text}"
    return text


def render_docx(template_bytes: bytes, context: dict, autoescape: bool | None = None) -> bytes:
    """Return the filled .docx as bytes, or raise RenderError."""
    if autoescape is None:
        autoescape = settings.LETTER_STUDIO["AUTOESCAPE"]
    try:
        tpl = DocxTemplate(io.BytesIO(template_bytes))
        # opening the parts early turns a broken file into a clear message
        tpl.init_docx()
    except Exception as exc:  # noqa: BLE001 - any failure here means "not a usable .docx"
        raise RenderError("That file isn't a .docx Word can open.", kind="file") from exc

    env = make_env(autoescape)
    try:
        tpl.render(context or {}, jinja_env=env, autoescape=autoescape)
    except TemplateError as exc:
        lines = [re.sub(r"\s+", " ", x).strip() for x in getattr(exc, "docx_context", []) or []]
        lines = [x for x in lines if x][:7]
        tag = ""
        for line in lines:
            m = re.search(r"\{[{%].*?[}%]\}", line)
            if m:
                tag = m.group(0)
                break
        part = getattr(getattr(tpl, "current_rendering_part", None), "partname", "") or ""
        raise RenderError(_friendly(exc), kind=type(exc).__name__, tag=tag, context=lines, part=str(part)) from exc
    except Exception as exc:  # noqa: BLE001 - filters or data can fail in many ways
        raise RenderError(f"The letter couldn't be filled: {exc}", kind=type(exc).__name__) from exc

    out = io.BytesIO()
    tpl.save(out)
    return out.getvalue()


def docx_text(docx_bytes: bytes) -> dict:
    """Plain text of the filled letter (body paragraphs, table rows, headers, footers) for quick checks."""
    from docx import Document

    doc = Document(io.BytesIO(docx_bytes))
    body = [p.text for p in doc.paragraphs]
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    headers, footers = [], []
    for section in doc.sections:
        headers += [p.text for p in section.header.paragraphs]
        footers += [p.text for p in section.footer.paragraphs]
    return {"body": body, "tables": tables, "headers": headers, "footers": footers}


# ---------------------------------------------------------------- validate

TAG_RE = re.compile(r"\{\{.*?\}\}|\{%.*?%\}")
FOOTNOTES_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
ENDNOTES_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"
MAX_EMPLOYEES = 5000


def _open(template_bytes):
    try:
        tpl = DocxTemplate(io.BytesIO(template_bytes))
        tpl.init_docx()
        return tpl
    except Exception as exc:  # noqa: BLE001 - any failure here means "not a usable .docx"
        raise RenderError("That file isn't a .docx Word can open.", kind="file") from exc


def template_parts(tpl):
    """[(path, label, xml)] for the body, headers, footers and notes, the way docxtpl fills them.

    The xml has one paragraph per line (as in docxtpl's render), so a Jinja line number
    points at a paragraph: line 1 is what comes before the first paragraph, line 2 the first one.
    """
    raw = [("word/document.xml", "Body", tpl.patch_xml(tpl.get_xml()))]
    for uri, label in ((tpl.HEADER_URI, "Header"), (tpl.FOOTER_URI, "Footer")):
        for _, part in tpl.get_headers_footers(uri):
            raw.append((str(part.partname).lstrip("/"), label, tpl.patch_xml(tpl.get_part_xml(part))))
    for part in tpl.docx.part.package.parts:
        if part.content_type in (FOOTNOTES_TYPE, ENDNOTES_TYPE):
            blob = part.blob.decode("utf-8") if isinstance(part.blob, bytes) else part.blob
            label = "Footnotes" if part.content_type == FOOTNOTES_TYPE else "Endnotes"
            raw.append((str(part.partname).lstrip("/"), label, tpl.patch_xml(blob)))
    return [(path, label, re.sub(r"<w:p([ >])", r"\n<w:p\1", xml)) for path, label, xml in raw]


def _line_info(xml, lineno, message=""):
    """Where a Jinja line number lands: paragraph index in the part, its text, the tag on it."""
    lines = xml.splitlines()
    if not lineno or lineno < 1 or lineno > len(lines):
        return {}
    text = re.sub(r"<[^>]+>", "", lines[lineno - 1])
    text = re.sub(r"\s+", " ", text).strip()
    tags = TAG_RE.findall(text)
    # the tag the message is about: the one holding a word the message quotes ('iff', '=', 'old')
    words = [w for w in re.findall(r"'([^']+)'", message or "") if w not in ("end of statement block", "end of print statement")]
    tag = next((t for t in tags if any(w in t for w in words)), tags[0] if tags else "")
    return {"line": lineno, "para_index": lineno - 2 if lineno >= 2 else None,
            "line_text": text[:300], "tag": tag[:200]}


def _template_lineno(exc):
    """The template line a runtime error happened on (Jinja rewrites tracebacks to point at it)."""
    tb, line = exc.__traceback__, None
    while tb is not None:
        if tb.tb_frame.f_code.co_filename == "<template>":
            line = tb.tb_lineno
        tb = tb.tb_next
    return line


def _undefined_name(text):
    m = re.search(r"'([^']+)' is undefined", text) or re.search(r"has no attribute '([^']+)'", text)
    return m.group(1) if m else ""


def validate_docx(template_bytes: bytes, employees=None, autoescape: bool | None = None) -> dict:
    """Check a tagged .docx with real Jinja, without making any file.

    1. Syntax: every part is parsed the way docxtpl parses it; a typing mistake comes back
       with the paragraph it is in.
    2. Employees (optional): the letter is filled for each employee's values; any that stop
       with an error, or that would give a file Word can't open, are listed.
    """
    from jinja2 import meta
    from lxml import etree

    if autoescape is None:
        autoescape = settings.LETTER_STUDIO["AUTOESCAPE"]
    tpl = _open(template_bytes)
    env = make_env(autoescape)
    parts = template_parts(tpl)

    syntax, compiled, variables = [], [], set()
    for path, label, xml in parts:
        try:
            ast = env.parse(xml)
            template = env.from_string(xml)  # also catches unknown filters such as | floot
        except TemplateSyntaxError as exc:
            where = _line_info(xml, exc.lineno, exc.message or "")
            if "Unexpected end of template" in (exc.message or ""):
                # Jinja only notices at the very end; the open block is somewhere above
                where = {"line": None, "para_index": None, "line_text": "", "tag": ""}
            friendly = _friendly(exc)
            for t in TAG_RE.findall(where.get("line_text") or ""):
                inner = re.sub(r"\\.", "", t)
                if inner.count('"') % 2 or inner.count("'") % 2:
                    friendly, where["tag"] = f"A quote isn't closed in {t[:80]}. Every \" or ' needs its pair.", t[:200]
                    break
            syntax.append({"part": path, "part_label": label, "message": exc.message or str(exc),
                           "friendly": friendly, **where})
            continue
        variables |= meta.find_undeclared_variables(ast)
        compiled.append((path, label, xml, template))

    result = {"ok": not syntax, "syntax": syntax, "variables": sorted(variables - {"currentDatetime"}),
              "parts": [{"part": p, "label": lbl} for p, lbl, _ in parts], "employees": None}
    if syntax or not employees:
        return result

    tested, failed = 0, []
    for k, emp in enumerate(employees[:MAX_EMPLOYEES]):
        label = str(emp.get("label") or f"Employee {k + 1}")
        values = emp.get("values") if isinstance(emp.get("values"), dict) else {}
        tested += 1
        for path, part_label, xml, template in compiled:
            try:
                out = template.render(values)
            except Exception as exc:  # noqa: BLE001 - data can break a template in many ways
                text = str(exc)
                failed.append({"index": k, "label": label, "part": path, "part_label": part_label,
                               "kind": type(exc).__name__, "message": text[:300], "friendly": _friendly(exc),
                               "name": _undefined_name(text), **_line_info(xml, _template_lineno(exc), text)})
                break
            try:
                etree.fromstring(out.encode("utf-8"))
            except etree.XMLSyntaxError as exc:
                failed.append({"index": k, "label": label, "part": path, "part_label": part_label,
                               "kind": "BrokenFile", "message": str(exc)[:300],
                               "friendly": "The filled letter would be a file Word can't open. This usually means a value has &, < or > and the tags print it unescaped, or a condition cuts a paragraph or table in half.",
                               "name": ""})
                break
    result["employees"] = {"tested": tested, "failed": failed, "limit": MAX_EMPLOYEES,
                           "skipped": max(0, len(employees) - MAX_EMPLOYEES)}
    result["ok"] = not failed
    return result
