"""Makes the made-up version sets in tests_ui/fixtures/combine_sets/ (python tests_ui/fixtures/make_combine_sets.py).

Each set is one letter in several versions, the way clients send them, for "Combine several versions". Every company, name and
amount is invented. The sets have different shapes so the combine is tested on more than one client:

  two/     2 versions: a few words differ inside a paragraph, one extra paragraph.
  grid/    4 versions that are 2 independent choices (promotion x bonus): must give 2 yes / no fields, not 4.
  cables/  8 versions like a real annual-review pack: letterhead in its own section, merge fields, a salary table where one row
           is only in some versions (a blank line right before the table), a promotion paragraph, a note in 2 versions, a
           consultant version with other words in 2 sentences and without the annexure and everything after it, an incentive
           section only in some versions (page break, numbered list, a small table, a signature picture) with a few extra words
           in some of those, a closing section, header and footer.
  shapes/  6 versions: one paragraph in 3 versions (if / elif / else, each with its own formatting), two shared paragraphs in
           another order in one version, a table with other columns in one version, an extra annexure in one version, a
           missing closing section in another, a few words that differ in one version.
"""
import base64
import io
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

OUT = Path(__file__).resolve().parent / "combine_sets"
# 1x1 pixel PNGs (a "logo" and a "signature"), so pictures in the versions can be checked
LOGO = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGNgYPj/HwADAgH/eL9GtQAAAABJRU5ErkJggg==")
SIGN = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGNg+M/wHwAEAQH/7yMK/gAAAABJRU5ErkJggg==")


def merge_field(p, name):
    """A Word MERGEFIELD showing «name», as mail-merge letters have them."""
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), f" MERGEFIELD {name} ")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = f"«{name}»"
    r.append(t)
    f.append(r)
    p._p.append(f)


def para(d, *parts, bold=False, align=None, italic=False, style=None):
    """A paragraph from text and ("field", name) parts."""
    p = d.add_paragraph(style=style)
    for x in parts:
        if isinstance(x, tuple):
            merge_field(p, x[1])
        else:
            r = p.add_run(x)
            r.bold, r.italic = bold or None, italic or None
    if align:
        p.alignment = align
    return p


def picture(d, png, width=0.5):
    d.add_paragraph().add_run().add_picture(io.BytesIO(png), width=Inches(width))


def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def numbered(d, text):
    p = d.add_paragraph(text, style="List Number")
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", "0"), ("w:numId", "1")):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        numPr.append(e)
    pPr.append(numPr)


def table(d, rows, bold_first=False):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.paragraphs[0].text = ""
            for x in cell if isinstance(cell, list) else [cell]:
                if isinstance(x, tuple):
                    merge_field(c.paragraphs[0], x[1])
                else:
                    c.paragraphs[0].add_run(x).bold = (bold_first and i == 0) or None
    return t


def save(d, folder, name):
    (OUT / folder).mkdir(parents=True, exist_ok=True)
    d.save(OUT / folder / f"{name}.docx")


def header_footer(d, head, foot):
    s = d.sections[0]
    s.header.paragraphs[0].text = head
    s.footer.paragraphs[0].text = foot


# ------------------------------------------------------------------------------------------------ two
def make_two():
    for name, promo in (("standard", False), ("promotion", True)):
        d = Document()
        header_footer(d, "Bluefin Foods - Private", "Bluefin Foods, 4 Harbour Lane")
        para(d, "Dear Arjun Mehta,")
        para(d, "Your new pay is INR 9,40,000 a year, effective 1 April 2026"
             + (", together with your promotion to Senior Engineer." if promo else "."))
        if promo:
            para(d, "Your new manager will be in touch about your new role.")
        para(d, "Regards,")
        para(d, "People Team")
        save(d, "two", name)


# ------------------------------------------------------------------------------------------------ grid
def make_grid():
    for name, promo, bonus in (("base", False, False), ("promotion", True, False), ("bonus", False, True), ("promotion-bonus", True, True)):
        d = Document()
        header_footer(d, "Tallow Paints - Confidential", "Tallow Paints, 9 Mill Road")
        para(d, "Dear Nisha Kulkarni,")
        para(d, "Thank you for a good year. Your review is complete.")
        if promo:
            para(d, "We are pleased to promote you to Team Lead from 1 April 2026.")
        rows = [["Item", "Amount (INR)"], ["Fixed pay", [("f", "Fixed_Pay")]]]
        if bonus:
            rows.append(["One-time bonus", [("f", "Bonus")]])
        rows.append(["Total", [("f", "Total")]])
        table(d, rows, bold_first=True)
        if bonus:
            para(d, "The bonus is paid with the April salary.")
        para(d, "Regards,")
        para(d, "HR Team")
        save(d, "grid", name)


# ------------------------------------------------------------------------------------------------ cables
CABLES = {
    # name: promo, insurance row, incentive section, the extra words in it, note, consultant
    "standard-insurance": dict(promo=False, ins=True, pli=False, change=False, note=False, consultant=False),
    "insurance-promotion": dict(promo=True, ins=True, pli=False, change=False, note=False, consultant=False),
    "pay-merged": dict(promo=False, ins=False, pli=False, change=False, note=True, consultant=False),
    "promotion-pay-merged": dict(promo=True, ins=False, pli=False, change=False, note=True, consultant=False),
    "incentive-new": dict(promo=False, ins=False, pli=True, change=False, note=False, consultant=False),
    "incentive-change": dict(promo=False, ins=False, pli=True, change=True, note=False, consultant=False),
    "incentive-change-promotion": dict(promo=True, ins=False, pli=True, change=True, note=False, consultant=False),
    "consultant": dict(promo=False, ins=False, pli=False, change=False, note=False, consultant=True),
}


def letter_head(d):
    picture(d, LOGO, 1.0)
    para(d, "Larkspur Cables Ltd", bold=True)
    para(d, "21 Canal Street, Pune 411001")
    d.add_section(WD_SECTION.CONTINUOUS)


def signature(d):
    picture(d, SIGN, 0.8)
    para(d, "Kavita Rao")
    para(d, "Head of People")


def make_cables():
    for name, c in CABLES.items():
        d = Document()
        header_footer(d, "Larkspur Cables Ltd - Confidential", "Larkspur Cables Ltd, 21 Canal Street, Pune")
        letter_head(d)
        para(d, "Date: 1 April 2026")
        para(d)
        para(d, ("f", "Emp_Name"))
        para(d, ("f", "Emp_Code"))
        para(d)
        para(d, "Subject: Annual Review and " + ("Retainer Fee" if c["consultant"] else "Compensation") + " Update for 2026", bold=True)
        para(d)
        para(d, "Dear ", ("f", "First_Name"), ",")
        para(d)
        para(d, "Thank you for your work this year. The company grew well and your part in that matters to us.")
        para(d)
        if c["consultant"]:
            para(d, "In line with our pay for performance approach, we are pleased to revise your retainer fee from INR ",
                 ("f", "Old_Fee"), "/- to INR ", ("f", "New_Fee"), "/-, effective ", ("f", "Eff_Date"), ".")
        else:
            para(d, "In line with our pay for performance approach, we are pleased to revise your compensation from INR ",
                 ("f", "Old_Total"), "/- to INR ", ("f", "New_Total"), "/-, effective ", ("f", "Eff_Date"), ".")
        para(d)
        if c["promo"]:
            para(d, "We are also pleased to promote you to ", ("f", "New_Title"), " from ", ("f", "Promo_Date"), ".")
            para(d)
        para(d, "For details of your pay, please see Annexure I. Write to your HR partner with any questions.")
        para(d)
        para(d, "Best regards,")
        signature(d)
        if not c["consultant"]:
            para(d)
            page_break(d)
            para(d)   # a blank line right before the table
            rows = [["Annexure I", "", ""], ["Item", "Last year", "This year"],
                    ["Basic", [("f", "Basic_Old")], [("f", "Basic_New")]], ["HRA", [("f", "HRA_Old")], [("f", "HRA_New")]]]
            if c["ins"]:
                rows.append(["Insurance", [("f", "Ins_Old")], [("f", "Ins_New")]])
            rows.append(["Total", [("f", "Old_Total")], [("f", "New_Total")]])
            table(d, rows, bold_first=True)
            para(d)
            para(d, "Your acceptance of this letter means you accept all its terms.")
            signature(d)
            para(d)
            para(d, "Please keep this letter confidential.")
            if c["note"]:
                para(d)
                para(d, "Note: Your incentive has been merged into your fixed pay from the date above.", italic=True)
            if c["pli"]:
                page_break(d)
                para(d, "Date: 1 April 2026")
                para(d, ("f", "Emp_Name"))
                para(d, "Subject: Incentive Plan for 2026", bold=True)
                para(d)
                para(d, "Dear ", ("f", "First_Name"), ",")
                para(d)
                if c["change"]:
                    para(d, "Your incentive target for 2026 is INR ", ("f", "New_Target"), "/- (as against INR ", ("f", "Old_Target"), "/- last year).")
                else:
                    para(d, "Your incentive target for 2026 is INR ", ("f", "New_Target"), "/-.")
                para(d)
                para(d, "The payout depends on:")
                numbered(d, "Company results")
                numbered(d, "Team results")
                numbered(d, "Your own rating")
                table(d, [["Rating", "Payout"], ["3 and above", "100%"]], bold_first=True)
                para(d)
                signature(d)
            page_break(d)
            para(d, "About the new wage rules", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            para(d)
            para(d, "The new wage rules change how some parts of pay are counted.")
            para(d, "Your total pay does not change because of them.")
        save(d, "cables", name)


# ------------------------------------------------------------------------------------------------ shapes
SHAPES = {
    "a-steady": dict(alt=0, swap=False, cols3=True, annex=False, closing=True, reply="within two weeks"),
    "b-steady-short": dict(alt=0, swap=False, cols3=True, annex=False, closing=False, reply="within two weeks"),
    "c-raise-reordered": dict(alt=1, swap=True, cols3=True, annex=False, closing=True, reply="within two weeks"),
    "d-raise-date": dict(alt=1, swap=False, cols3=True, annex=False, closing=True, reply="by 30 April 2026"),
    "e-award-relocation": dict(alt=2, swap=False, cols3=True, annex=True, closing=True, reply="within two weeks"),
    "f-award-simple-table": dict(alt=2, swap=False, cols3=False, annex=False, closing=True, reply="within two weeks"),
}


def make_shapes():
    alts = [
        ("We have kept your pay at its current level this year.", dict(align=WD_ALIGN_PARAGRAPH.CENTER)),
        ("Your pay goes up by 6 percent from the first of April.", dict(bold=True)),
        ("You will receive a one-time award of INR 25,000 in May.", dict(italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT)),
    ]
    for name, c in SHAPES.items():
        d = Document()
        header_footer(d, "Quarry Lane Books", "Quarry Lane Books, 2 Station Road")
        para(d, "Dear Ravi Shetty,")
        para(d, "This letter sums up your yearly review.")
        para(d)
        text, fmt = alts[c["alt"]]
        para(d, text, **fmt)
        para(d)
        p3, p4 = "Your benefits stay the same as last year.", "Your leave balance carries over to next year."
        for t in ((p4, p3) if c["swap"] else (p3, p4)):
            para(d, t)
        para(d)
        if c["cols3"]:
            table(d, [["Item", "Old", "New"], ["Pay", "6,00,000", "6,36,000"], ["Total", "6,00,000", "6,36,000"]], bold_first=True)
        else:
            table(d, [["What", "Amount"], ["Award", "25,000"]], bold_first=True)
        para(d)
        para(d, f"Please sign a copy of this letter and send it back to HR {c['reply']}.")
        para(d, "Kind regards,")
        para(d, "Leela Menon")
        if c["annex"]:
            page_break(d)
            para(d, "Annexure II - Relocation", bold=True)
            para(d, "We will pay for your move to the new office.")
            para(d, "Keep the receipts and send them to HR.")
        if c["closing"]:
            para(d)
            para(d, "About your data", bold=True)
            para(d, "We keep your review for three years.")
            para(d, "You can ask HR for a copy at any time.")
        save(d, "shapes", name)


if __name__ == "__main__":
    if OUT.exists():
        shutil.rmtree(OUT)
    make_two()
    make_grid()
    make_cables()
    make_shapes()
    print("wrote", ", ".join(sorted(p.name for p in OUT.iterdir())))
