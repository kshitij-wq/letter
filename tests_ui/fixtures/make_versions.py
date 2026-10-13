"""Makes the four made-up letter versions in tests_ui/fixtures/versions/ (python tests_ui/fixtures/make_versions.py).

Same header and footer everywhere. Differences: the paragraph under the subject (page 1), an extra paragraph only in
"promotion", a closing paragraph missing in "no-increase", and one table row. "variable-pay" and "bonus-only" share their
paragraph. All names and amounts are made up.
"""
from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parent / "versions"

CASES = {
    "promotion": {
        "para": "We are pleased to confirm your promotion to the next grade, effective 1 April 2026.",
        "extra": "Your new title is Senior Analyst.",
        "sign": True,
        "row": ("New fixed pay", "1,320,000"),
    },
    "variable-pay": {
        "para": "Your variable pay for the year has been set at 12 percent of your annual target.",
        "extra": None,
        "sign": True,
        "row": ("Variable pay", "150,000"),
    },
    "no-increase": {
        "para": "After this year's review your pay stays at its current level.",
        "extra": None,
        "sign": False,
        "row": ("Fixed pay", "1,200,000"),
    },
    "bonus-only": {
        "para": "Your variable pay for the year has been set at 12 percent of your annual target.",
        "extra": None,
        "sign": True,
        "row": ("One-time bonus", "50,000"),
    },
}


def make(name, c):
    d = Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Northwind Appliances - Confidential"
    sec.footer.paragraphs[0].text = "Northwind Appliances, 12 Example Road"
    d.add_paragraph("Date: 1 April 2026")
    d.add_paragraph("Dear Asha Rao,")
    d.add_paragraph("Subject: Annual compensation review")
    d.add_paragraph(c["para"])
    if c["extra"]:
        d.add_paragraph(c["extra"])
    d.add_paragraph("This letter is confidential and is meant only for you.")
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Amount (INR)"
    t.cell(1, 0).text, t.cell(1, 1).text = c["row"]
    d.add_paragraph("All amounts are before tax.")
    if c["sign"]:
        d.add_paragraph("Please sign and return a copy of this letter within seven days.")
    d.add_paragraph("Regards,")
    d.add_paragraph("HR Team")
    d.save(OUT / f"{name}.docx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for n, c in CASES.items():
        make(n, c)
    print("wrote", ", ".join(CASES))
