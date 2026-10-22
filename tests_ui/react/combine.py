"""Combine several versions, end to end, with the real page and real Jinja.

    python tests_ui/react/combine.py        (set CHROMIUM_PATH if Playwright has no browser of its own)

Round 1 drops the four made-up versions in tests_ui/fixtures/versions/ on the home page: the screen (versions, the yes / no fields
with their suggested names, the differences), renaming a field, combining, undo / redo, the preview in the page and the download.
Then every made-up set in tests_ui/fixtures/combine_sets/ (made by make_combine_sets.py: 2, 4, 6 and 8 versions of different shapes)
goes through the same screen. Each download is checked with renderer.engine.validate_docx and filled with docxtpl once per version
(every yes / no field "yes" for exactly the versions the screen lists for it); the result must be that version: the same text
paragraph for paragraph and row for row, and the same pictures, page breaks, sections, numbered paragraphs and paragraph / first-run
formatting. Merge fields the page turned into tags are compared as fields («…»), whatever their values.
"""
import io
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "letterstudio.settings")
import django  # noqa: E402

django.setup()

from docx import Document  # noqa: E402
from docxtpl import DocxTemplate  # noqa: E402
from jinja2 import Undefined  # noqa: E402
from playwright.sync_api import expect, sync_playwright  # noqa: E402

from renderer import engine  # noqa: E402
from renderer.engine import validate_docx  # noqa: E402

sys.path.insert(0, str(ROOT / "tests_ui"))
from common import FIX, launch, start_static  # noqa: E402
VERS = FIX / "versions"
SETS = FIX / "combine_sets"
NAMES = ["promotion", "variable-pay", "no-increase", "bonus-only"]
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
errors = []


class AnyField(Undefined):
    """A tag whose value isn't in the test data prints as «F», like the merge field it came from."""
    def __str__(self):
        return "«F»"

    def __float__(self):
        return 0.0

    def __int__(self):
        return 0

    def __getattr__(self, name):
        if name.startswith("__"):
            raise AttributeError(name)
        return AnyField()

    def __getitem__(self, key):
        return AnyField()


def fill(tagged, values):
    """docxtpl with the page's Jinja environment, the yes / no fields set and every other tag printed as «F»."""
    tpl = DocxTemplate(io.BytesIO(tagged))
    env = engine.make_env(True)
    env.undefined = AnyField
    for k in ("indian_comma", "int_comma", "format_date"):
        f = env.filters[k]
        env.filters[k] = (lambda f: lambda v, *a, **kw: "«F»" if isinstance(v, Undefined) else f(v, *a, **kw))(f)
    tpl.render({"custom_schema": values}, jinja_env=env, autoescape=True)
    out = io.BytesIO()
    tpl.save(out)
    return out.getvalue()


def fields(s):
    return re.sub(r"«F»(\s*«F»)+", "«F»", re.sub(r"«[^»]*»|<<[^>]*>>", "«F»", s))


def ptext(p):
    return "".join(x.text or "" for x in p.iter(W + "t"))


def doc_lines(source):
    """The text of a .docx in order: every non-empty paragraph, and a table as one line per row."""
    d = Document(source)
    out = []
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            t = ptext(el).strip()
            if t:
                out.append(fields(t))
        elif tag == "tbl":
            for tr in el.iter(W + "tr"):
                out.append("| " + " | ".join(fields(" ".join(ptext(p) for p in tc.iter(W + "p")).strip()) for tc in tr.findall(W + "tc")))
    return out


def doc_shape(source):
    """Everything else that must survive: pictures, page breaks, sections, numbered paragraphs, and per text paragraph its style,
    alignment and first-run bold / italic."""
    d = Document(source)
    body = d.element.body
    fmt = []
    for p in body.iter(W + "p"):
        t = ptext(p).strip()
        if not t:
            continue
        ppr = p.find(W + "pPr")
        val = lambda path: (ppr.find(path).get(W + "val") if ppr is not None and ppr.find(path) is not None else None)  # noqa: E731
        r = next((r for r in p.iter(W + "r") if (r.findtext(W + "t") or "").strip()), None)
        rpr = r.find(W + "rPr") if r is not None else None
        b = rpr is not None and rpr.find(W + "b") is not None
        i = rpr is not None and rpr.find(W + "i") is not None
        fmt.append((fields(t)[:40], val(W + "pStyle"), val(W + "jc"), b, i))
    return {
        "pictures": len(list(body.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))),
        "page breaks": sum(1 for x in body.iter(W + "br") if x.get(W + "type") == "page"),
        "sections": len(d.sections),
        "numbered": sum(1 for x in body.iter(W + "numPr")),
        "format": fmt,
    }


def blanks(source):
    return sum(1 for p in Document(source).element.body.iter(W + "p") if not ptext(p).strip() and p.find(f".//{W}drawing") is None)


def check_versions(tagged, flags, files, where, blank_slack=0):
    """Fill the download once per version and compare with that version."""
    res = validate_docx(tagged)
    assert res["ok"], (where, res["syntax"])
    for f in files:
        label = f.stem
        values = {fl["field"].split(".", 1)[1]: ("yes" if label in fl["versions"] else "no") for fl in flags}
        out = io.BytesIO(fill(tagged, values))
        got, want = doc_lines(out), doc_lines(f)
        assert got == want, f"{where}, {label}: filled letter differs from the version\n got : {got}\n want: {want}"
        gs, ws = doc_shape(out), doc_shape(f)
        for k in ws:
            assert gs[k] == ws[k], f"{where}, {label}: {k} differ\n got : {gs[k]}\n want: {ws[k]}"
        gb, wb = blanks(out), blanks(f)
        assert abs(gb - wb) <= blank_slack, f"{where}, {label}: {gb} blank lines, the version has {wb}"


def download(pg):
    pg.get_by_test_id("step-export").click()
    with pg.expect_download() as dl:
        pg.click('button:has-text("Download tagged .docx")')
        expect(pg.get_by_role("dialog")).to_contain_text("Download anyway?")   # the yes / no fields are not in any data yet: the gate asks
        pg.get_by_test_id("confirm-ok").click()
    return Path(dl.value.path()).read_bytes()


def combine_files(pg, url, files):
    pg.goto(url)
    with pg.expect_file_chooser() as fc:
        pg.get_by_test_id("combine-entry").click()
    fc.value.set_files([str(f) for f in files])
    expect(pg.get_by_test_id("combine-screen")).to_be_visible()
    expect(pg.get_by_test_id("combine-go")).to_be_enabled()
    flags = pg.evaluate("window.__lt.combineView().flags")
    pg.get_by_test_id("combine-go").click()
    expect(pg.get_by_test_id("combine-screen")).to_have_count(0)
    return flags


def main():
    srv, _ = start_static()
    with sync_playwright() as p:
        b = launch(p)
        pg = b.new_context(viewport={"width": 1440, "height": 900}, accept_downloads=True).new_page()
        pg.on("pageerror", lambda e: errors.append(str(e)))
        pg.on("console", lambda m: errors.append(m.text) if m.type == "error" and "404" not in m.text else None)
        url = f"http://127.0.0.1:{srv.server_address[1]}/"
        pg.goto(url)

        # --- a file that is not a .docx: plain error on the home page, no combine screen
        with pg.expect_file_chooser() as fc:
            pg.get_by_test_id("combine-entry").click()
        fc.value.set_files(str(FIX / "emps.csv"))
        expect(pg.get_by_role("alert")).to_contain_text("emps.csv is not a .docx file.")
        expect(pg.get_by_test_id("combine-screen")).to_have_count(0)

        # --- one version only: the screen opens with a hint and Combine stays off
        with pg.expect_file_chooser() as fc:
            pg.get_by_test_id("combine-entry").click()
        fc.value.set_files(str(VERS / "promotion.docx"))
        expect(pg.get_by_test_id("combine-screen")).to_be_visible()
        expect(pg.get_by_test_id("combine-hint")).to_contain_text("Add at least one more version")
        expect(pg.get_by_test_id("combine-go")).to_be_disabled()

        # a bad file next to good ones is named on the screen, the good ones are kept
        with pg.expect_file_chooser() as fc:
            pg.get_by_test_id("combine-add").click()
        fc.value.set_files([str(VERS / f"{n}.docx") for n in NAMES[1:]] + [str(FIX / "emps.csv")])
        expect(pg.get_by_test_id("combine-errors")).to_contain_text("emps.csv is not a .docx file.")
        for i, n in enumerate(NAMES):
            expect(pg.get_by_test_id(f"combine-label-{i}")).to_have_value(n)
        expect(pg.get_by_test_id("combine-go")).to_be_enabled()   # nothing is required
        expect(pg.get_by_test_id("combine-problems")).to_have_count(0)

        # --- the summary, the yes / no fields, the differences
        expect(pg.get_by_test_id("combine-summary")).to_contain_text("8 parts are the same in every version")
        expect(pg.get_by_test_id("combine-summary")).to_contain_text("3 places differ")
        boxes = pg.locator("[data-testid^=combine-flag-name-]")
        names = [boxes.nth(i).input_value() for i in range(boxes.count())]
        flags = pg.evaluate("window.__lt.combineView().flags")
        assert "promotion" in [f["name"] for f in flags], flags
        assert names == [f["name"] for f in flags], (names, flags)
        k = [f["name"] for f in flags].index("promotion")
        expect(pg.get_by_test_id(f"combine-flag-{k}")).to_contain_text("“yes” for: promotion")
        region0 = pg.get_by_test_id("combine-region-0")
        expect(region0).to_contain_text("We are pleased to confirm your promotion")
        expect(region0).to_contain_text("When")
        expect(pg.get_by_role("button", name="Set who gets which version (optional)")).to_have_count(0)
        expect(pg.get_by_test_id("combine-screen")).not_to_contain_text("letter_version")
        # rename a field: the differences read the new name at once
        box = pg.get_by_test_id(f"combine-flag-name-{k}")
        box.fill("promoted")
        box.press("Enter")
        expect(region0).to_contain_text("promoted is “yes”")
        # reorder with the keyboard-reachable arrows: the base follows its version, the name stays
        pg.get_by_role("button", name="Move bonus-only up").click()
        expect(pg.get_by_test_id("combine-label-2")).to_have_value("bonus-only")
        pg.get_by_role("button", name="Move bonus-only down").click()
        expect(pg.get_by_test_id("combine-label-3")).to_have_value("bonus-only")
        expect(pg.get_by_test_id("combine-base-0")).to_be_checked()
        expect(pg.get_by_test_id(f"combine-flag-name-{k}")).to_have_value("promoted")
        flags = pg.evaluate("window.__lt.combineView().flags")

        # --- combine
        pg.get_by_test_id("combine-go").click()
        expect(pg.get_by_test_id("combine-screen")).to_have_count(0)
        expect(pg.get_by_test_id("step-logic")).to_have_attribute("aria-selected", "true")
        expect(pg.get_by_text("Combined 4 versions. Each difference is an if / else on a yes / no field").first).to_be_visible()
        assert pg.evaluate("window.__lt.HISTORY.undo.length") == 1, "combining must be one undo step"
        expect(pg.get_by_test_id("placeholder-field-note")).to_contain_text("custom_schema.promoted")
        n_logic = pg.evaluate("window.__lt.state.logic.length")

        # --- Validate says in plain words which fields to add, and points back to Conditions
        pg.get_by_test_id("step-check").click()
        expect(pg.get_by_text("Choose which employee fields decide the version").first).to_be_visible()
        pg.get_by_role("button", name="Open conditions").first.click()
        expect(pg.get_by_test_id("step-logic")).to_have_attribute("aria-selected", "true")

        tagged = download(pg)
        check_versions(tagged, flags, [VERS / f"{n}.docx" for n in NAMES], "round 1")
        # nobody's field says yes: only the text that is the same everywhere, and the "else" paragraph
        other = doc_lines(io.BytesIO(fill(tagged, {f["field"].split(".", 1)[1]: "no" for f in flags})))
        assert any("Asha Rao" in x for x in other) and not any("variable pay" in x for x in other), other

        # --- undo restores the base letter in one step; redo brings the combined template back
        pg.get_by_role("button", name="Undo").first.click()
        pg.wait_for_timeout(300)
        assert pg.evaluate("window.__lt.state.logic.length") == 0, "undo left logic lines"
        pg.evaluate("window.__lt.build()")
        xml = pg.evaluate("window.__lt.state.built.xml['word/document.xml']")
        assert "After this year" not in xml and "We are pleased to confirm your promotion" in xml
        pg.get_by_role("button", name="Redo").first.click()
        pg.wait_for_timeout(300)
        assert pg.evaluate("window.__lt.state.logic.length") == n_logic

        # --- the preview in the page, once per version
        pg.click('button:has-text("Open employee preview")')
        pg.get_by_role("button", name="Show and change the values used in this preview").click()
        want = {
            "promotion": ("We are pleased to confirm your promotion", "Your new title is Senior Analyst.", "Please sign and return"),
            "variable-pay": ("Your variable pay for the year", "12 percent", "Please sign and return"),
            "no-increase": ("After this year's review", "stays at its current level", "Regards"),
            "bonus-only": ("Your variable pay for the year", "One-time bonus", "Please sign and return"),
        }
        for n in NAMES:
            for f in flags:
                box = pg.get_by_label(f"Value of {f['field']}")
                box.fill("yes" if n in f["versions"] else "no")
                box.press("Enter")
            pg.wait_for_timeout(500)
            text = pg.locator("#paper").inner_text()
            for w in want[n]:
                assert w in text, f"preview for {n} lacks {w!r}"
            if n == "no-increase":
                assert "Please sign and return" not in text and "promotion to the next grade" not in text, "preview for no-increase shows other versions' text"
        pg.get_by_role("button", name="Back to tagged template").click()

        # --- round 2: another base version, the same result
        pg.goto(url)
        with pg.expect_file_chooser() as fc:
            pg.get_by_test_id("combine-entry").click()
        fc.value.set_files([str(VERS / f"{n}.docx") for n in NAMES])
        expect(pg.get_by_test_id("combine-screen")).to_be_visible()
        pg.get_by_test_id("combine-base-2").check()
        flags = pg.evaluate("window.__lt.combineView().flags")
        pg.get_by_test_id("combine-go").click()
        expect(pg.get_by_test_id("combine-screen")).to_have_count(0)
        check_versions(download(pg), flags, [VERS / f"{n}.docx" for n in NAMES], "round 2 (base no-increase)")

        # --- every made-up set: 2, 4, 6 and 8 versions of different shapes
        # blank lines: an if / else between two whole tables needs a paragraph of its own (Word joins two tables that touch),
        # and a blank line that some versions have in place of another version's text goes with the nearest difference
        for name, slack in (("two", 0), ("grid", 0), ("shapes", 1), ("cables", 1)):
            files = sorted((SETS / name).glob("*.docx"))
            flags = combine_files(pg, url, files)
            check_versions(download(pg), flags, files, name, slack)
        b.close()
    if errors:
        print("PAGE ERRORS:", *errors, sep="\n  ")
        sys.exit(1)
    print("combine test passed")


if __name__ == "__main__":
    main()
