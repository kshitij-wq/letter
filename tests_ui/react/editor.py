"""The in-app Word editor (Edit letter view), driven through the real page and the real Django back end.

    python tests_ui/react/editor.py        (set CHROMIUM_PATH when Playwright has no browser of its own)

Types into SuperDoc, saves, and checks what the engine has afterwards: every tag still there, the paper shows the edit, real Jinja
(Django + docxtpl) still validates, the downloaded .docx reopens in python-docx and renders for an employee from emps.csv. Also: undo
of an editor change, the blocked save when a tag is deleted, a tag split over runs, an undefined field, a missing endif, no data,
and header/footer kept. Fails on any page error. Build the page first: python tools/build_frontend.py
"""
from __future__ import annotations

import asyncio
import io
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import validate as V  # noqa: E402  (the helpers: server, page setup, ok())
from playwright.async_api import async_playwright, expect  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
FIX = V.FIX
ok = V.ok
CSV = (FIX / "emps.csv").read_text()

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "letterstudio.settings")
sys.path.insert(0, str(ROOT))
import django  # noqa: E402

django.setup()
from renderer.engine import render_docx, validate_docx  # noqa: E402
import docx as pydocx  # noqa: E402


async def open_letter(page, name: str, data: bool = False):
    await page.goto(page.url)
    await page.wait_for_function("window.__lt && window.__ltUi")
    async with page.expect_file_chooser() as fc:
        await page.get_by_text("Choose a .docx file").click()
    await (await fc.value).set_files(str(FIX / name))
    # a letter that was just opened starts on Tags & fields with the client's letter; Edit letter is one click away
    await page.wait_for_function("window.__lt.state.view === 'source' && window.__lt.state.tab === 'tags'")
    await page.get_by_test_id("seg-edit").click()
    await page.wait_for_function("window.__lt.state.view === 'edit'")
    await page.wait_for_function("window.__sdReady !== false && !!window.__sd", timeout=20000)
    await page.wait_for_timeout(1200)   # the editor settles its own document
    if data:
        await page.evaluate("t => window.__lt.ingestCsv(t, 'emps.csv')", CSV)


def run(page, text: str):
    return page.locator(".superdoc-text-run", has_text=text).first


async def type_at_end(page, anchor: str, text: str):
    await run(page, anchor).click()
    await page.keyboard.press("End")
    await page.keyboard.type(text)


async def select_first(page, anchor: str, n: int):
    """Select the first n characters of the paragraph (a double click does not select words in the editor)."""
    await run(page, anchor).click()
    await page.keyboard.press("Home")
    for _ in range(n):
        await page.keyboard.press("Shift+ArrowRight")


async def save(page):
    btn = page.get_by_test_id("editor-save")
    try:
        await btn.click(timeout=1500)   # the 2-second auto-save may get there first
    except Exception:
        pass
    await expect(page.get_by_test_id("editor-status")).to_have_text("Saved in this browser", timeout=15000)


async def paras(page):
    return await page.evaluate("window.__lt.state.an.paras.map(p => p.model.text)")


async def download(page) -> bytes:
    async with page.expect_download() as dl:
        await page.get_by_role("button", name="Download tagged .docx").click()
        try: await page.get_by_test_id("confirm-ok").click(timeout=3000)   # open items: "Download anyway?"
        except Exception: pass
    path = await (await dl.value).path()
    return Path(path).read_bytes()


async def tags_in(page):
    return await page.evaluate("window.__lt.docTags(window.__lt.state.srcBuf)")


async def main():
    proc, url = V.start_django()
    try:
        async with async_playwright() as p:
            browser = await V.launch(p, args=["--no-sandbox"])
            ctx, page = await V.new_page(browser, url)
            page.set_default_timeout(20000)

            # ---- 1. conditions letter + data: type, bold, save, downloads, conditions still render --------------------
            await open_letter(page, "conditions_tagged.docx", data=True)
            ok(await page.get_by_test_id("editor").count() == 1, "Edit letter opens from the view switch")
            ok(await page.locator(".tag-chip-block").count() >= 5 and await page.locator(".tag-chip-var").count() >= 1, "tags are drawn as chips (yellow and indigo)")
            before_tags = await tags_in(page)
            await type_at_end(page, "Regards,", " warmly")
            await page.wait_for_function("['unsaved','saving','saved'].includes(window.__lt.ed.status)", timeout=5000)
            await select_first(page, "Welcome to", 7)      # a word, then bold
            await page.keyboard.press("Control+b")
            await save(page)
            ok(any("Regards, warmly" in t for t in await paras(page)), "the engine's letter shows the typed text")
            await expect(page.get_by_text("Edited in Edit letter").first).to_be_visible()
            ok(await page.get_by_text("Not changed yet").count() == 0, "the header no longer says Not changed yet after an Edit letter save")
            ok(await tags_in(page) == before_tags, "every tag is still in the letter after the save")
            await page.get_by_role("tab", name="Validate").click()
            await expect(page.get_by_text("The full check passed")).to_be_visible(timeout=20000)
            ok(True, "Validate still runs real Jinja on the edited letter")
            await page.get_by_role("tab", name="Document").click()
            blob = await download(page)
            ok(validate_docx(blob)["syntax"] == [], "the downloaded template passes real Jinja (docxtpl)")
            names = [p.text for p in pydocx.Document(io.BytesIO(blob)).paragraphs]
            ok(any("Regards, warmly" in n for n in names), "the downloaded .docx reopens in python-docx with the edit")
            yes = render_docx(blob, {"first_name": "Asha", "promotion": "yes", "dept": "Finance"})
            no = render_docx(blob, {"first_name": "Ravi", "promotion": "no", "dept": "R&D"})
            t_yes = "\n".join(p.text for p in pydocx.Document(io.BytesIO(yes)).paragraphs)
            t_no = "\n".join(p.text for p in pydocx.Document(io.BytesIO(no)).paragraphs)
            ok("You are promoted." in t_yes and "You are promoted." not in t_no, "a conditional paragraph still renders by employee")
            ok("research" in t_no and "research" not in t_yes, "an inline condition still renders by employee")

            # ---- 2. undo one editor change ------------------------------------------------------------------------
            await page.evaluate("window.__lt.setView('edit')")
            await page.get_by_test_id("editor-undo").click()
            await page.wait_for_function("window.__lt.state.an.paras.every(p => !p.model.text.includes('Regards, warmly'))")
            ok(True, "Undo editor change brings the original letter back")
            await ctx.close()

            # ---- 3. table cell, header/footer, a tag split over runs ------------------------------------------------
            ctx, page = await V.new_page(browser, url)
            page.set_default_timeout(20000)
            await open_letter(page, "with_header.docx")
            await run(page, "Basic").click()
            await page.keyboard.press("Tab")
            await page.keyboard.type("45,000")
            await save(page)
            ok(any("45,000" in t for t in await paras(page)), "a typed table cell value reaches the engine's letter")
            parts = await page.evaluate("window.__lt.state.parts.map(p => p.kind)")
            ok("header" in parts, "the header part is kept")
            hdr = await page.evaluate("window.__lt.state.an.paras.filter(p => p.kind === 'header').map(p => p.model.text).join(' ')")
            ok("Confidential" in hdr, "the header text is kept")
            blob = await download(page)
            ok(any("Confidential" in p.text for s in pydocx.Document(io.BytesIO(blob)).sections for p in s.header.paragraphs), "the downloaded .docx still has its header")
            await ctx.close()

            ctx, page = await V.new_page(browser, url)
            page.set_default_timeout(20000)
            await open_letter(page, "acme_tagged.docx", data=True)
            tags0 = await tags_in(page)
            ok(len(tags0) >= 30, "a letter with tags in table cells opens with its tags")
            await select_first(page, "Following the annual", 9)   # bold the first word of a paragraph
            await page.keyboard.press("Control+b")
            await run(page, "Dear").click()
            await page.keyboard.press("Home")
            for _ in range(8):
                await page.keyboard.press("ArrowRight")
            for _ in range(5):
                await page.keyboard.press("Shift+ArrowRight")   # five characters inside {{ first_name }}: italic splits the tag's runs
            await page.keyboard.press("Control+i")
            await save(page)
            ok(await tags_in(page) == tags0, "a tag split over runs by formatting still counts as the same tag")
            blob = await download(page)
            ok(validate_docx(blob)["syntax"] == [], "and the template still passes real Jinja")
            filled = render_docx(blob, {"first_name": "Asha", "last_name": "Menon", "ctc": 1140000, "basic": 540000, "basic_mo": 45000, "hra": 1, "hra_mo": 1, "employee_schema": {"updated_designation": "Lead", "department": "Fin"}, "current_date": "1 April 2026"})
            ok("Dear Asha Menon" in "\n".join(p.text for p in pydocx.Document(io.BytesIO(filled)).paragraphs), "an employee letter renders from the edited template")

            # ---- 4. deleting a tag is blocked; restoring it ---------------------------------------------------------
            await run(page, "Your department is").click()
            await page.keyboard.press("End")
            for _ in range(len("{{ employee_schema.department }}.") + 1):
                await page.keyboard.press("Backspace")
            await expect(page.get_by_test_id("editor-blocked")).to_be_visible(timeout=15000)   # auto-save after 2 s tries and is blocked
            ok("employee_schema.department" in await page.get_by_test_id("editor-blocked").inner_text(), "a deleted tag blocks the save and is named")
            ok(await tags_in(page) == tags0, "nothing was saved while it was blocked")
            await page.get_by_test_id("editor-restore-tags").click()
            await expect(page.get_by_test_id("editor-blocked")).to_have_count(0)
            await page.wait_for_timeout(1500)
            ok(any("employee_schema.department" in t for t in await paras(page)), "Restore those tags puts the letter back")

            # ---- 5. an undefined field and a missing endif ------------------------------------------------------------
            await type_at_end(page, "HR Team", " {{ nope_field }}")
            await save(page)
            await page.get_by_role("tab", name="Validate").click()
            await expect(page.get_by_text("nope_field").first).to_be_visible(timeout=20000)
            ok(True, "an undefined field typed in the editor shows up in Validate")
            await page.get_by_role("tab", name="Document").click()
            await page.evaluate("window.__lt.setView('edit')")
            await page.wait_for_timeout(1500)
            await type_at_end(page, "HR Team", " {% if x %}")
            await save(page)
            await page.get_by_role("tab", name="Validate").click()
            await expect(page.get_by_text("The full check found")).to_be_visible(timeout=20000)
            ok(True, "a missing endif typed in the editor fails the real Jinja check")
            await ctx.close()

            # ---- 6. no data file, select words -> Tag this text ----------------------------------------------------------
            ctx, page = await V.new_page(browser, url)
            page.set_default_timeout(20000)
            await open_letter(page, "with_header.docx")
            await select_first(page, "Your department is", 14)
            await expect(page.get_by_test_id("editor-tag-text")).to_be_visible()
            await page.get_by_test_id("editor-tag-text").click()
            await page.wait_for_function("window.__lt.state.view === 'source'")
            ok(await page.evaluate("!!window.__lt.state.sel.range"), "Tag this text selects the same words in the client letter")
            ok(await page.get_by_role("complementary").count() >= 1, "the Selected text panel opens")
            await ctx.close()

            await browser.close()
    finally:
        proc.kill()
    if V.PROBLEMS:
        print("\n".join(["", "PROBLEMS:"] + V.PROBLEMS))
        sys.exit(1)
    print("\nEditor test passed.")


if __name__ == "__main__":
    asyncio.run(main())
